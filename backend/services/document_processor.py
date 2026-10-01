"""
Procesador de documentos - Extrae texto de PDF y DOCX, divide en chunks
"""
from pathlib import Path
import PyPDF2
from docx import Document
import logging
import re
import shutil
import subprocess
import tempfile
from typing import Optional, List, Dict, Any

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """
    Procesa documentos PDF y DOCX, extrae texto y lo divide en chunks
    """

    # Objetivo aproximado en tokens/palabras. Se prefiere mantener párrafos y oraciones completas.
    CHUNK_TARGET_MIN_TOKENS = 120
    CHUNK_TARGET_MAX_TOKENS = 180
    CHUNK_HARD_MAX_TOKENS = 220
    CHUNK_OVERLAP_TOKENS = 30
    MIN_CHUNK_CHARS = 50

    @staticmethod
    def process_document(filepath: str, filename: str) -> Dict[str, Any]:
        """
        Procesa un documento y extrae chunks
        Args:
            filepath: Ruta del archivo
            filename: Nombre del archivo
        Returns:
            Diccionario con información del documento
        """
        filepath_obj = Path(filepath)
        filetype = filepath_obj.suffix.lower()

        logger.info(f"Procesando {filename} ({filetype})")

        try:
            if filetype == '.pdf':
                result = DocumentProcessor._process_pdf(filepath_obj)
            elif filetype == '.docx':
                result = DocumentProcessor._process_docx(filepath_obj)
            elif filetype == '.doc':
                result = DocumentProcessor._process_doc(filepath_obj)
            else:
                raise ValueError(f"Tipo de archivo no soportado: {filetype}")

            result['filename'] = filename
            result['filetype'] = filetype
            page_texts = result.pop('page_texts', None)
            if page_texts:
                result['chunks'] = DocumentProcessor._create_chunks_from_pages(page_texts)
            else:
                result['chunks'] = DocumentProcessor._create_chunks(result['text'])
            logger.info(f"✓ {filename}: {result['pages']} páginas, {len(result['chunks'])} chunks")

            return result

        except Exception as e:
            logger.error(f"Error procesando {filename}: {str(e)}")
            raise

    @staticmethod
    def _process_pdf(filepath: Path) -> Dict[str, Any]:
        """
        Extrae texto de un archivo PDF
        Args:
            filepath: Ruta del archivo PDF
        Returns:
            Diccionario con texto y número de páginas
        """
        try:
            text_content: List[str] = []
            num_pages = 0

            page_texts: List[tuple] = []

            with open(filepath, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                num_pages = len(pdf_reader.pages)

                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        text = page.extract_text()
                        if text:
                            cleaned = DocumentProcessor._clean_text(text, preserve_paragraphs=True)
                            if cleaned:
                                text_content.append(cleaned)
                                page_texts.append((page_num + 1, cleaned))
                    except Exception as e:
                        logger.warning(f"Error extrayendo página {page_num + 1}: {str(e)}")

            full_text = '\n'.join(text_content)

            return {
                'text': full_text,
                'pages': num_pages,
                'page_texts': page_texts
            }

        except Exception as e:
            logger.error(f"Error procesando PDF: {str(e)}")
            raise

    @staticmethod
    def _process_docx(filepath: Path) -> Dict[str, Any]:
        """
        Extrae texto de un archivo DOCX
        Args:
            filepath: Ruta del archivo DOCX 
        Returns:
            Diccionario con texto y número de páginas
        """
        try:
            doc = Document(filepath)
            text_content: List[str] = []

            for para in doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)

            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content.append(cell.text)

            full_text = '\n\n'.join(text_content)
            full_text = DocumentProcessor._clean_text(full_text, preserve_paragraphs=True)
            word_count = len(full_text.split())
            estimated_pages = max(1, word_count // 250)

            return {
                'text': full_text,
                'pages': estimated_pages
            }
        except Exception as e:
            logger.error(f"Error procesando DOCX: {str(e)}")
            raise

    @staticmethod
    def _process_doc(filepath: Path) -> Dict[str, Any]:
        """
        Extrae texto de un archivo DOC clásico usando Word COM y, si no está disponible,
        una conversión a DOCX con LibreOffice.

        Args:
            filepath: Ruta del archivo DOC

        Returns:
            Diccionario con texto y número estimado de páginas
        """
        try:
            return DocumentProcessor._process_doc_with_win32(filepath)
        except Exception as com_error:
            logger.warning(
                f"No se pudo procesar DOC con Word COM ({com_error}). Intentando fallback a DOCX..."
            )

        try:
            return DocumentProcessor._process_doc_via_conversion(filepath)
        except Exception as conversion_error:
            logger.error(f"Error procesando DOC con fallback: {conversion_error}")
            raise RuntimeError(
                "No se pudo procesar el archivo .doc. Necesita Microsoft Word o LibreOffice para convertirlo. "
                "Comprueba que el entorno tenga Word/LibreOffice instalado y accesible."
            ) from conversion_error

    @staticmethod
    def _process_doc_with_win32(filepath: Path) -> Dict[str, Any]:
        """Procesa DOC usando Word COM en Windows."""
        word_app = None
        document = None

        try:
            try:
                import win32com.client as win32
            except ImportError as e:
                raise RuntimeError(
                    "No se puede procesar archivos .doc porque falta win32com.client"
                ) from e

            word_app = win32.DispatchEx("Word.Application")
            word_app.Visible = False
            word_app.DisplayAlerts = 0

            document = word_app.Documents.Open(str(filepath), ReadOnly=True)

            text_content: List[str] = []

            for paragraph in document.Paragraphs:
                paragraph_text = paragraph.Range.Text.strip()
                if paragraph_text:
                    text_content.append(paragraph_text)

            if not text_content:
                raw_text = document.Content.Text
                if raw_text:
                    text_content.append(raw_text)

            full_text = '\n\n'.join(text_content)
            full_text = DocumentProcessor._clean_text(full_text, preserve_paragraphs=True)

            try:
                estimated_pages = int(document.ComputeStatistics(2))
            except Exception:
                word_count = len(full_text.split())
                estimated_pages = max(1, word_count // 250)

            return {
                'text': full_text,
                'pages': estimated_pages
            }

        finally:
            try:
                if document is not None:
                    document.Close(False)
            except Exception:
                pass

            try:
                if word_app is not None:
                    word_app.Quit()
            except Exception:
                pass

    @staticmethod
    def _process_doc_via_conversion(filepath: Path) -> Dict[str, Any]:
        """Convierte DOC a DOCX y procesa el resultado con python-docx."""
        soffice = shutil.which("soffice") or shutil.which("soffice.exe")
        if not soffice:
            raise RuntimeError("No hay LibreOffice/soffice disponible para convertir DOC a DOCX")

        with tempfile.TemporaryDirectory(prefix="doc_convert_") as tmpdir:
            output_dir = Path(tmpdir)
            command = [
                soffice,
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                str(output_dir),
                str(filepath),
            ]
            result = subprocess.run(command, capture_output=True, text=True)
            if result.returncode != 0:
                raise RuntimeError(
                    f"La conversión falló: {result.stderr.strip() or result.stdout.strip() or 'sin detalle'}"
                )

            converted_path = output_dir / f"{filepath.stem}.docx"
            if not converted_path.exists():
                converted_candidates = list(output_dir.glob("*.docx"))
                if not converted_candidates:
                    raise RuntimeError("No se generó ningún DOCX durante la conversión")
                converted_path = converted_candidates[0]

            converted_doc = Document(converted_path)
            text_content: List[str] = []
            for para in converted_doc.paragraphs:
                if para.text.strip():
                    text_content.append(para.text)

            for table in converted_doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_content.append(cell.text)

            full_text = '\n\n'.join(text_content)
            full_text = DocumentProcessor._clean_text(full_text, preserve_paragraphs=True)
            word_count = len(full_text.split())
            estimated_pages = max(1, word_count // 250)

            return {
                'text': full_text,
                'pages': estimated_pages
            }

    @staticmethod
    def _clean_text(text: str, preserve_paragraphs: bool = False) -> str:
        """
        Limpia el texto extraído
        
        Args:
            text: Texto a limpiar
            preserve_paragraphs: Si se conservan los saltos de párrafo
            
        Returns:
            Texto limpio
        """
        text = text.replace('\r\n', '\n').replace('\r', '\n')
        text = text.replace('\x00', '')
        text = text.replace('\x0c', '')

        if preserve_paragraphs:
            text = re.sub(r'[ \t]+', ' ', text)
            text = re.sub(r'\n[ \t]*\n+', '\n\n', text)
            text = re.sub(r'(?<!\n)\n(?!\n)', ' ', text)
        else:
            text = re.sub(r'\s+', ' ', text)

        text = text.strip()
        return text

    @staticmethod
    def _create_chunks(
        text: str,
        chunk_size: Optional[int] = None,
        overlap: Optional[int] = None,
        page: Optional[int] = None
    ) -> List[Dict[str, Any]]:
        """Divide texto en chunks semánticos por párrafos y oraciones."""
        if not text:
            return []

        target_min_tokens = DocumentProcessor.CHUNK_TARGET_MIN_TOKENS if chunk_size is None else chunk_size
        target_max_tokens = DocumentProcessor.CHUNK_TARGET_MAX_TOKENS
        overlap_tokens = DocumentProcessor.CHUNK_OVERLAP_TOKENS if overlap is None else overlap
        overlap_tokens = max(20, min(40, overlap_tokens))
        hard_max_tokens = max(target_max_tokens, DocumentProcessor.CHUNK_HARD_MAX_TOKENS)

        paragraphs = [paragraph.strip() for paragraph in re.split(r'\n{2,}', text) if paragraph.strip()]
        if not paragraphs:
            paragraphs = [text.strip()]

        chunks: List[Dict[str, Any]] = []

        current_parts: List[str] = []
        current_tokens = 0
        token_tail: List[str] = []

        def flush_current() -> None:
            nonlocal current_parts, current_tokens, token_tail
            if not current_parts:
                return

            content = '\n\n'.join(current_parts).strip()
            if len(content) > DocumentProcessor.MIN_CHUNK_CHARS:
                chunks.append({'content': content, 'page': page})

                # Mantener cola corta de tokens para dar solapamiento entre chunks.
                content_tokens = content.split()
                token_tail = content_tokens[-overlap_tokens:] if overlap_tokens > 0 else []
            else:
                token_tail = []

            current_parts = []
            current_tokens = 0

        def add_segment(segment: str, segment_tokens: int) -> None:
            nonlocal current_parts, current_tokens

            if not segment:
                return

            if current_parts and current_tokens + segment_tokens > hard_max_tokens:
                flush_current()

            current_parts.append(segment.strip())
            current_tokens += segment_tokens

            if current_tokens >= target_max_tokens:
                flush_current()

                if token_tail:
                    carry = ' '.join(token_tail).strip()
                    if carry:
                        current_parts = [carry]
                        current_tokens = len(token_tail)

        for paragraph in paragraphs:
            paragraph_tokens = DocumentProcessor._estimate_tokens(paragraph)
            if paragraph_tokens > hard_max_tokens:
                for sentence in DocumentProcessor._split_into_sentences(paragraph):
                    sentence_tokens = DocumentProcessor._estimate_tokens(sentence)
                    if sentence_tokens > hard_max_tokens:
                        for fragment in DocumentProcessor._split_long_text(sentence, hard_max_tokens):
                            add_segment(fragment, DocumentProcessor._estimate_tokens(fragment))
                    else:
                        add_segment(sentence, sentence_tokens)
                continue

            if current_parts and current_tokens >= target_min_tokens and current_tokens + paragraph_tokens > target_max_tokens:
                flush_current()

            add_segment(paragraph, paragraph_tokens)

        flush_current()

        logger.debug(
            f"Creados {len(chunks)} chunks semánticos (~{target_min_tokens}-{target_max_tokens} tokens, overlap={overlap_tokens})"
        )
        return chunks

    @staticmethod
    def _create_chunks_from_pages(page_texts: List[tuple]) -> List[Dict[str, Any]]:
        """Parte cada página por separado para conservar el número de página."""
        chunks: List[Dict[str, Any]] = []
        for page_number, page_text in page_texts:
            chunks.extend(
                DocumentProcessor._create_chunks(page_text, page=page_number)
            )
        return chunks

    @staticmethod
    def _estimate_tokens(text: str) -> int:
        """Aproxima tokens con palabras separadas por espacio."""
        return len(text.split()) if text else 0

    @staticmethod
    def _split_into_sentences(paragraph: str) -> List[str]:
        """Divide un párrafo en oraciones conservando el orden."""
        paragraph = paragraph.strip()
        if not paragraph:
            return []

        sentences = [sentence.strip() for sentence in re.split(r'(?<=[.!?])\s+', paragraph) if sentence.strip()]
        return sentences if sentences else [paragraph]

    @staticmethod
    def _split_long_text(text: str, max_tokens: int) -> List[str]:
        """Divide un texto demasiado largo en fragmentos por palabras sin perder demasiado contexto."""
        words = text.split()
        if not words:
            return []

        fragments: List[str] = []
        start = 0
        overlap = max(20, min(40, max_tokens // 6))

        while start < len(words):
            end = min(start + max_tokens, len(words))
            fragment = ' '.join(words[start:end]).strip()
            if fragment:
                fragments.append(fragment)

            if end >= len(words):
                break

            start = max(end - overlap, start + 1)

        return fragments

    @staticmethod
    def validate_file(filepath: str) -> bool:
        """
        Valida que el archivo sea legible
        
        Args:
            filepath: Ruta del archivo
            
        Returns:
            True si el archivo es válido
        """
        try:
            filepath_obj = Path(filepath)
            if not filepath_obj.exists():
                logger.error(f"Archivo no existe: {filepath}")
                return False

            filetype = filepath_obj.suffix.lower()
            if filetype not in ['.pdf', '.docx', '.doc']:
                logger.error(f"Tipo de archivo no soportado: {filetype}")
                return False

            return True
        except Exception as e:
            logger.error(f"Error validando archivo: {str(e)}")
            return False